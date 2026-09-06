"""
Conciliación de movimientos bancarios de Demo TestFire.

Este módulo compara los movimientos extraídos desde la página web con
un libro maestro local llamado ``conciliacion_master.xlsx``.

El archivo maestro debe contener las columnas:

- ID Transacción
- Fecha/Hora
- ID Cuenta
- Acción
- Monto

La clase no crea ni modifica el maestro. Solamente lo lee y devuelve el
resultado de la comparación para que otros módulos puedan generar el
reporte Excel y registrar la auditoría.
"""

from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class ConciliadorMovimientos:
    """
    Compara los movimientos de TestFire con un archivo maestro local.

    La conciliación utiliza el ID de transacción como identificador principal.
    Si el ID existe en ambas fuentes, también compara la fecha, la cuenta,
    la acción y el monto.
    """

    # Columnas que obligatoriamente debe contener la primera fila del maestro.
    ENCABEZADOS_REQUERIDOS = (
        "ID Transacción",
        "Fecha/Hora",
        "ID Cuenta",
        "Acción",
        "Monto",
    )

    def __init__(self, ruta_maestro: str | Path | None = None):
        """
        Configura la ubicación del libro maestro.

        Args:
            ruta_maestro:
                Ruta opcional del archivo que se desea utilizar. Si no se
                entrega, se usa automáticamente:
                ``data/entrada/conciliacion_master.xlsx``.
        """

        # __file__ representa src/conciliacion.py. Sus dos carpetas superiores
        # permiten encontrar la raíz del proyecto sin depender de la ubicación
        # desde la cual se ejecute el comando "py main.py".
        raiz_proyecto = Path(__file__).resolve().parent.parent

        if ruta_maestro is None:
            self.ruta_maestro = (
                raiz_proyecto
                / "data"
                / "entrada"
                / "conciliacion_master.xlsx"
            )
        else:
            self.ruta_maestro = Path(ruta_maestro).resolve()

    @staticmethod
    def _obtener_valor(movimiento: dict[str, Any], *claves: str) -> Any:
        """
        Obtiene un dato desde un movimiento aceptando nombres alternativos.

        Esto permite trabajar tanto con claves internas como
        ``id_transaccion`` como con encabezados visibles del Excel, por
        ejemplo ``ID Transacción``.

        Args:
            movimiento:
                Diccionario que contiene una transacción.

            *claves:
                Posibles nombres bajo los cuales puede encontrarse el dato.

        Returns:
            Any:
                Primer valor encontrado. Retorna None si ninguna clave existe.
        """

        for clave in claves:
            if clave in movimiento:
                return movimiento[clave]

        return None

    @staticmethod
    def _normalizar_id(valor: Any) -> str:
        """
        Convierte un identificador a texto comparable.

        Excel puede leer un número como ``17434.0`` mientras Selenium lo
        entrega como ``"17434"``. Esta normalización evita considerarlos
        distintos cuando representan el mismo identificador.
        """

        if valor is None:
            return ""

        texto = str(valor).strip()

        if texto.endswith(".0") and texto[:-2].isdigit():
            texto = texto[:-2]

        return texto

    @staticmethod
    def _normalizar_texto(valor: Any) -> str:
        """Convierte un valor a texto sin espacios sobrantes y en minúsculas."""

        if valor is None:
            return ""

        return " ".join(str(valor).split()).casefold()

    @staticmethod
    def _normalizar_fecha(valor: Any) -> str:
        """
        Convierte fechas de Excel o texto al formato ``AAAA-MM-DD HH:MM``.

        Si el valor no coincide con los formatos habituales, se mantiene
        como texto limpio para que igualmente pueda compararse.
        """

        if valor is None:
            return ""

        if isinstance(valor, datetime):
            return valor.strftime("%Y-%m-%d %H:%M")

        texto = " ".join(str(valor).split())

        formatos = (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%d-%m-%Y %H:%M:%S",
            "%d-%m-%Y %H:%M",
        )

        for formato in formatos:
            try:
                fecha = datetime.strptime(texto, formato)
                return fecha.strftime("%Y-%m-%d %H:%M")
            except ValueError:
                continue

        return texto.casefold()

    @staticmethod
    def _normalizar_monto(valor: Any) -> Decimal:
        """
        Convierte montos numéricos o textos monetarios a Decimal.

        Acepta ejemplos como ``1234``, ``$1234.00``, ``$1.234,00`` y
        ``-$1,234.00``. Decimal evita errores comunes de aproximación de float.

        Raises:
            ValueError:
                Cuando el monto está vacío o no puede convertirse.
        """

        if valor is None or str(valor).strip() == "":
            raise ValueError("Se encontró un movimiento sin monto.")

        if isinstance(valor, Decimal):
            return valor.quantize(Decimal("0.01"))

        if isinstance(valor, (int, float)):
            return Decimal(str(valor)).quantize(Decimal("0.01"))

        texto = str(valor).strip()
        texto = texto.replace("$", "").replace(" ", "")

        # También se aceptan montos negativos escritos entre paréntesis.
        if texto.startswith("(") and texto.endswith(")"):
            texto = f"-{texto[1:-1]}"

        # Cuando aparecen punto y coma, el último separador se considera
        # decimal y el otro se considera separador de miles.
        if "," in texto and "." in texto:
            if texto.rfind(",") > texto.rfind("."):
                texto = texto.replace(".", "").replace(",", ".")
            else:
                texto = texto.replace(",", "")
        elif "," in texto:
            # Una coma seguida por dos dígitos se interpreta como decimal.
            parte_decimal = texto.rsplit(",", maxsplit=1)[-1]
            if len(parte_decimal) == 2:
                texto = texto.replace(".", "").replace(",", ".")
            else:
                texto = texto.replace(",", "")

        try:
            return Decimal(texto).quantize(Decimal("0.01"))
        except InvalidOperation as error:
            raise ValueError(
                f"El monto '{valor}' no tiene un formato válido."
            ) from error

    def _normalizar_movimiento(self, movimiento: dict[str, Any]) -> dict[str, Any]:
        """
        Transforma un movimiento web o del maestro a una estructura común.

        Returns:
            dict[str, Any]:
                Movimiento con las claves internas utilizadas durante la
                conciliación.
        """

        id_transaccion = self._obtener_valor(
            movimiento,
            "id_transaccion",
            "ID Transacción",
        )

        if self._normalizar_id(id_transaccion) == "":
            raise ValueError(
                "Se encontró un movimiento sin ID de transacción."
            )

        monto_original = self._obtener_valor(
            movimiento,
            "monto",
            "Monto",
        )

        return {
            "id_transaccion": self._normalizar_id(id_transaccion),
            "fecha_hora": self._normalizar_fecha(
                self._obtener_valor(
                    movimiento,
                    "fecha_hora",
                    "Fecha/Hora",
                )
            ),
            "id_cuenta": self._normalizar_id(
                self._obtener_valor(
                    movimiento,
                    "id_cuenta",
                    "ID Cuenta",
                )
            ),
            "accion": self._normalizar_texto(
                self._obtener_valor(
                    movimiento,
                    "accion",
                    "Acción",
                )
            ),
            "monto": self._normalizar_monto(monto_original),
        }

    def _leer_maestro(self) -> list[dict[str, Any]]:
        """
        Lee y valida los movimientos de ``conciliacion_master.xlsx``.

        Returns:
            list[dict[str, Any]]:
                Registros del archivo maestro con sus valores normalizados.

        Raises:
            FileNotFoundError:
                Cuando el archivo maestro no existe.

            ValueError:
                Cuando el libro está vacío, le faltan columnas o contiene
                IDs duplicados.
        """

        if not self.ruta_maestro.exists():
            raise FileNotFoundError(
                "No se encontró el archivo maestro en: "
                f"{self.ruta_maestro}"
            )

        libro = load_workbook(
            filename=self.ruta_maestro,
            data_only=True,
            read_only=True,
        )

        try:
            hoja = libro["Movimientos"] if "Movimientos" in libro.sheetnames else libro.active

            encabezados = [
                str(celda.value).strip() if celda.value is not None else ""
                for celda in hoja[1]
            ]

            faltantes = [
                encabezado
                for encabezado in self.ENCABEZADOS_REQUERIDOS
                if encabezado not in encabezados
            ]

            if faltantes:
                raise ValueError(
                    "El archivo maestro no contiene las columnas: "
                    + ", ".join(faltantes)
                )

            posiciones = {
                encabezado: encabezados.index(encabezado)
                for encabezado in self.ENCABEZADOS_REQUERIDOS
            }

            movimientos_maestro = []
            ids_encontrados = set()

            for numero_fila, fila in enumerate(
                hoja.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                # Las filas completamente vacías no representan movimientos.
                if all(valor is None for valor in fila):
                    continue

                movimiento = {
                    encabezado: fila[indice]
                    for encabezado, indice in posiciones.items()
                }

                normalizado = self._normalizar_movimiento(movimiento)
                id_transaccion = normalizado["id_transaccion"]

                if id_transaccion in ids_encontrados:
                    raise ValueError(
                        "El archivo maestro contiene el ID duplicado "
                        f"'{id_transaccion}' en la fila {numero_fila}."
                    )

                ids_encontrados.add(id_transaccion)
                movimientos_maestro.append(normalizado)

            if not movimientos_maestro:
                raise ValueError(
                    "El archivo conciliacion_master.xlsx no contiene movimientos."
                )

            return movimientos_maestro
        finally:
            libro.close()

    @staticmethod
    def _buscar_diferencias(
        movimiento_web: dict[str, Any],
        movimiento_maestro: dict[str, Any],
    ) -> list[str]:
        """
        Devuelve los nombres de los campos que no coinciden.

        El ID no se incluye porque ya se utilizó para relacionar ambos
        movimientos.
        """

        nombres_campos = {
            "fecha_hora": "Fecha/Hora",
            "id_cuenta": "ID Cuenta",
            "accion": "Acción",
            "monto": "Monto",
        }

        return [
            nombre_visible
            for clave, nombre_visible in nombres_campos.items()
            if movimiento_web[clave] != movimiento_maestro[clave]
        ]

    def conciliar(
        self,
        movimientos: list[dict[str, Any]],
    ) -> tuple[list[dict[str, Any]], dict[str, int]]:
        """
        Compara los movimientos extraídos con el archivo maestro.

        Se generan cuatro situaciones posibles:

        - ``Conciliado``: existe en ambas fuentes y todos sus datos coinciden.
        - ``Con diferencias``: el ID existe, pero uno o más campos cambian.
        - ``Solo en TestFire``: el ID extraído no existe en el maestro.
        - ``Solo en maestro``: el ID del maestro no apareció en TestFire.

        Args:
            movimientos:
                Lista de diccionarios devuelta por TestFireBot.

        Returns:
            tuple[list[dict[str, Any]], dict[str, int]]:
                Resultados detallados y resumen numérico de la conciliación.
        """

        if not movimientos:
            raise ValueError(
                "No se recibieron movimientos de TestFire para conciliar."
            )

        print("[CONCILIACIÓN] Leyendo archivo maestro...")

        maestro = self._leer_maestro()
        web = [
            self._normalizar_movimiento(movimiento)
            for movimiento in movimientos
        ]

        maestro_por_id = {
            movimiento["id_transaccion"]: movimiento
            for movimiento in maestro
        }
        ids_web = set()
        resultados = []

        print("[CONCILIACIÓN] Comparando movimientos...")

        for movimiento_web in web:
            id_transaccion = movimiento_web["id_transaccion"]

            if id_transaccion in ids_web:
                raise ValueError(
                    "Los movimientos extraídos contienen el ID duplicado "
                    f"'{id_transaccion}'."
                )

            ids_web.add(id_transaccion)
            movimiento_maestro = maestro_por_id.get(id_transaccion)

            if movimiento_maestro is None:
                estado = "Solo en TestFire"
                detalle = "El ID no existe en el archivo maestro."
            else:
                diferencias = self._buscar_diferencias(
                    movimiento_web,
                    movimiento_maestro,
                )

                if diferencias:
                    estado = "Con diferencias"
                    detalle = "Campos diferentes: " + ", ".join(diferencias)
                else:
                    estado = "Conciliado"
                    detalle = "El movimiento coincide con el archivo maestro."

            resultados.append(
                {
                    **movimiento_web,
                    "estado_conciliacion": estado,
                    "detalle_conciliacion": detalle,
                }
            )

        # También identificamos movimientos esperados en el maestro que no
        # fueron entregados por TestFire durante la ejecución actual.
        for movimiento_maestro in maestro:
            if movimiento_maestro["id_transaccion"] in ids_web:
                continue

            resultados.append(
                {
                    **movimiento_maestro,
                    "estado_conciliacion": "Solo en maestro",
                    "detalle_conciliacion": (
                        "El ID del maestro no fue encontrado en TestFire."
                    ),
                }
            )

        resumen = {
            "total_testfire": len(web),
            "total_maestro": len(maestro),
            "conciliados": sum(
                resultado["estado_conciliacion"] == "Conciliado"
                for resultado in resultados
            ),
            "con_diferencias": sum(
                resultado["estado_conciliacion"] == "Con diferencias"
                for resultado in resultados
            ),
            "solo_testfire": sum(
                resultado["estado_conciliacion"] == "Solo en TestFire"
                for resultado in resultados
            ),
            "solo_maestro": sum(
                resultado["estado_conciliacion"] == "Solo en maestro"
                for resultado in resultados
            ),
        }

        resumen["no_conciliados"] = (
            resumen["con_diferencias"]
            + resumen["solo_testfire"]
            + resumen["solo_maestro"]
        )

        print(
            "[CONCILIACIÓN] Proceso terminado. "
            f"Conciliados: {resumen['conciliados']} | "
            f"No conciliados: {resumen['no_conciliados']}"
        )

        return resultados, resumen
